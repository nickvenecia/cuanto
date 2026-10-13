"""Spreadsheets: budget lines in from a CSV file, and a whole budget out to Excel.

CSV in. The header row is found by its words, wherever it is in the first rows, so a sheet laid out by hand
works: Description (or Item, Name), Amount / Qty, Units, Unit (or Per, Type), Rate (or Price, Cost), Total,
and, when it has them, Category, Account, Code, Section, Fringes, Notes. Without a Category column, a row
with words and no numbers is a heading: the lines under it make an account. A line with a total and no rate
comes in as an allowance for that total. Category codes in the thousands place a category the way most
charts of accounts do (1000s above the line, 2000s and 3000s production, 4000s post, the rest other).

Excel out: a Top Sheet, the Detail, the Globals and the Fringes, as values (numbers stay numbers; text that
starts like a formula stays text, so an exported file can't run anything).
"""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from cuanto.core import calc, money
from cuanto.core.model import (
    BASE_NAMES,
    KIND_NAMES,
    SECTION_NAMES,
    SECTIONS,
    SYMBOLS,
    UNIT_NAMES,
    Account,
    Budget,
    Category,
    Line,
)

MAX_BYTES = 5_000_000
MAX_ROWS = 20_000
HEADER_SCAN = 25  # the header row is somewhere in the first rows


class SheetError(ValueError):
    """A file Cuánto can't read as budget lines; the message says what to do."""


# ---------------------------------------------------------------------------- CSV in
COLUMNS = {  # our name -> the words a header might use (compared without spaces, case or punctuation)
    "category": ("category", "cat", "categoryname", "department", "dept"),
    "category_code": ("categorycode", "catcode", "catno", "cat#", "categoryno"),
    "account": ("account", "accountname", "acct", "acctname"),
    "code": ("code", "accountcode", "acctcode", "acct#", "account#", "accountno", "acctno", "#", "no"),
    "description": ("description", "desc", "item", "lineitem", "name", "details", "detail"),
    "amount": ("amount", "amt", "qty", "quantity", "count", "people", "x"),
    "units": ("units", "unitqty", "days", "time", "duration", "length"),
    "unit": ("unit", "unittype", "per", "type", "uom", "unitofmeasure"),
    "rate": ("rate", "price", "cost", "unitcost", "unitprice", "dayrate"),
    "total": ("total", "subtotal", "linetotal", "totalcost", "sum", "estimate"),
    "section": ("section", "atlbtl", "level", "aboveorbelow"),
    "fringes": ("fringes", "fringe"),
    "notes": ("notes", "note", "comments", "comment", "remarks"),
}
UNIT_WORDS = {
    "days": ("day", "days", "d", "dy", "dys"),
    "weeks": ("week", "weeks", "wk", "wks", "w"),
    "hours": ("hour", "hours", "hr", "hrs", "h"),
    "months": ("month", "months", "mo", "mos", "mth"),
    "flat": ("flat", "fl", "flatfee", "fee"),
    "allow": ("allow", "allowance", "alw", "allw", "lump", "lumpsum", "ls"),
    "each": ("each", "ea", "item", "items", "pc", "pcs", "piece"),
    "units": ("unit", "units", "u"),
}


def _key(text: str) -> str:
    return re.sub(r"[^a-z0-9#]", "", (text or "").casefold())


def unit_of(text: str) -> str:
    k = _key(text)
    for unit, words in UNIT_WORDS.items():
        if k in words:
            return unit
    return "allow"


def section_of(text: str, code: str = "") -> str:
    """A section from a word ("ATL", "Post"), else from a category code's thousands."""
    k = _key(text)
    if k in ("atl", "above", "abovetheline", "abovetheline"):
        return "atl"
    if k.startswith("post"):
        return "post"
    if k in ("other", "otherbtl", "general"):
        return "other"
    if k in ("production", "btl", "below", "belowtheline", "prod"):
        return "production"
    digits = re.match(r"\s*(\d{4})", code or "")
    if digits:
        first = digits.group(1)[0]
        return {"1": "atl", "2": "production", "3": "production", "4": "post"}.get(first, "other")
    return "production"


@dataclass
class Parsed:
    """What a CSV holds, before it goes into a budget."""

    categories: list[Category] = field(default_factory=list)
    lines: int = 0
    skipped: int = 0
    fringe_names: dict[str, list[str]] = field(default_factory=dict)  # line id -> fringe names from the sheet


def _rows(path: Path) -> list[list[str]]:
    if path.stat().st_size > MAX_BYTES:
        raise SheetError(f"{path.name} is too big for a budget. Split it, then import each part.")
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as e:
        raise SheetError(f"Cuánto can't read {path.name} as CSV.") from e
    if len(rows) > MAX_ROWS:
        raise SheetError(f"{path.name} has too many rows for a budget.")
    return rows


def _header(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:HEADER_SCAN]):
        found: dict[str, int] = {}
        for col, cell in enumerate(row):
            k = _key(cell)
            for ours, words in COLUMNS.items():
                if ours not in found and k in words:
                    found[ours] = col
                    break
        has_numbers = any(x in found for x in ("amount", "rate", "total", "units"))
        if "description" in found and has_numbers:
            # "Amount" means money when there's no Rate and no Total beside it
            if "amount" in found and "rate" not in found and "total" not in found:
                found["total"] = found.pop("amount")
            return i, found
    raise SheetError(
        "Cuánto couldn't find the column titles. Put Description and Rate (or Total) in one row."
    )


def _cell(row: list[str], cols: dict[str, int], name: str) -> str:
    i = cols.get(name)
    return row[i].strip()[:2000] if i is not None and i < len(row) else ""


def _number(text: str) -> str:
    """A cell as a number someone typed: "$1,250.00" -> "1250.00"; "(50)" -> "-50"; "" stays ""."""
    t = text.strip()
    if not t:
        return ""
    neg = t.startswith("(") and t.endswith(")")
    d = money.parse(t.strip("()"))
    if d == 0 and not re.search(r"\d", t):
        return ""
    return f"{-d if neg else d:f}"


def read_csv(path: Path) -> Parsed:
    """Budget lines from a CSV file. Raises SheetError (or OSError) with words for the person."""
    rows = _rows(path)
    head, cols = _header(rows)
    out = Parsed()
    cats: dict[str, Category] = {}
    current_cat: Category | None = None
    current_acc: Account | None = None
    by_category = "category" in cols

    def category(name: str, code: str, section_word: str) -> Category:
        key = (code or name).casefold()
        c = cats.get(key)
        if c is None:
            c = Category(code=code, name=name or code, section=section_of(section_word, code))
            cats[key] = c
            out.categories.append(c)
        return c

    def account(c: Category, name: str, code: str) -> Account:
        key = (code or name).casefold()
        a = next((a for a in c.accounts if (a.code or a.name).casefold() == key), None)
        if a is None:
            a = Account(code=code, name=name or code or c.name)
            c.accounts.append(a)
        return a

    for row in rows[head + 1 :]:
        if not any(cell.strip() for cell in row):
            continue
        desc = _cell(row, cols, "description")
        amount, units = _number(_cell(row, cols, "amount")), _number(_cell(row, cols, "units"))
        rate, total = _number(_cell(row, cols, "rate")), _number(_cell(row, cols, "total"))
        has_money = bool(rate or total)
        if by_category:
            cname, ccode = _cell(row, cols, "category"), _cell(row, cols, "category_code")
            if not (cname or ccode) and not has_money:
                out.skipped += 1
                continue
            c = (
                category(cname or "Imported", ccode, _cell(row, cols, "section"))
                if (cname or ccode)
                else (current_cat or category("Imported", "", ""))
            )
            aname, acode = _cell(row, cols, "account"), _cell(row, cols, "code")
            a = (
                account(c, aname or c.name, acode)
                if (aname or acode)
                else (current_acc if current_acc is not None and current_cat is c else account(c, c.name, ""))
            )
            current_cat, current_acc = c, a
        else:
            if desc and not has_money and not amount and not units:  # a heading: an account starts
                current_cat = current_cat or category("Imported", "", _cell(row, cols, "section"))
                current_acc = account(current_cat, desc, _cell(row, cols, "code"))
                continue
            if current_cat is None:
                current_cat = category("Imported", "", _cell(row, cols, "section"))
            if current_acc is None:
                current_acc = account(current_cat, "Imported Lines", "")
            a = current_acc
        if not desc and not has_money:
            out.skipped += 1
            continue
        unit_text = _cell(row, cols, "unit")
        if rate:
            ln = Line(
                description=desc, amount=amount or "1", units=units or "1", unit=unit_of(unit_text), rate=rate
            )
        else:
            ln = Line(description=desc, amount="1", units="1", unit="allow", rate=total)
        ln.notes = _cell(row, cols, "notes")
        names = [n.strip() for n in re.split(r"[,;/]", _cell(row, cols, "fringes")) if n.strip()]
        if names:
            out.fringe_names[ln.id] = names
        a.lines.append(ln)
        out.lines += 1
    if not out.lines:
        raise SheetError(f"There are no budget lines in {path.name}.")
    return out


def _same(a_code: str, a_name: str, b_code: str, b_name: str) -> bool:
    """The same category or account: the same code when both have one, else the same name."""
    if a_code.strip() and b_code.strip():
        return a_code.strip().casefold() == b_code.strip().casefold()
    return bool(a_name.strip()) and a_name.strip().casefold() == b_name.strip().casefold()


def merge_into(budget: Budget, parsed: Parsed) -> int:
    """Put what a CSV holds into a budget: categories and accounts with the same code (or name) are joined,
    new ones go at the end of their section. Fringe names from the sheet find the budget's fringes.
    Returns how many lines came in."""
    from cuanto.core import rates

    n = 0
    for pc in parsed.categories:
        mine = next((c for c in budget.categories if _same(c.code, c.name, pc.code, pc.name)), None)
        if mine is None:
            at = max((i + 1 for i, c in enumerate(budget.categories) if SECTIONS.index(c.section) <= SECTIONS.index(pc.section)), default=0)  # fmt: skip
            budget.categories.insert(at, pc)
            mine = pc
            for a in pc.accounts:
                for ln in a.lines:
                    ln.fringes = rates.fringe_ids(budget, parsed.fringe_names.get(ln.id, []))
                    n += 1
            continue
        for pa in pc.accounts:
            acc = next((a for a in mine.accounts if _same(a.code, a.name, pa.code, pa.name)), None)
            if acc is None:
                mine.accounts.append(pa)
                acc = pa
            else:
                acc.lines.extend(pa.lines)
            for ln in pa.lines:
                ln.fringes = rates.fringe_ids(budget, parsed.fringe_names.get(ln.id, []))
                n += 1
    return n


# ---------------------------------------------------------------------------- Excel out
def _safe(text: str) -> str:
    return (text or "").replace("\x00", "")


def write_xlsx(budget: Budget, dest: Path, profile_name: str = "") -> None:
    """The budget as an Excel workbook: Top Sheet, Detail, Globals, Fringes. Written beside `dest` first,
    then moved into place, so a failed export never leaves half a file."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    t = calc.totals(budget)
    sym = SYMBOLS.get(budget.currency, "")
    fmt = f'"{sym}"#,##0.00;[Red]-"{sym}"#,##0.00' if sym else "#,##0.00"
    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="C2600E")
    band = PatternFill("solid", fgColor="F6E7DB")

    wb = Workbook()

    def text(ws, row: int, col: int, value: str, font=None):
        c = ws.cell(row=row, column=col)
        c.value = _safe(value)
        c.data_type = "s"  # never a formula, whatever it starts with
        if font is not None:
            c.font = font
        return c

    def num(ws, row: int, col: int, value: Decimal, font=None, money_fmt: bool = True):
        c = ws.cell(row=row, column=col, value=float(value))
        c.number_format = fmt if money_fmt else "#,##0.##"
        c.alignment = Alignment(horizontal="right")
        if font is not None:
            c.font = font
        return c

    def heads(ws, row: int, titles: list[str]) -> None:
        for i, h in enumerate(titles, 1):
            c = text(ws, row, i, h, head_font)
            c.fill = head_fill

    # ---- top sheet
    ws = wb.active
    ws.title = "Top Sheet"
    text(ws, 1, 1, budget.name or "Untitled Budget", Font(bold=True, size=16))
    info = [
        x for x in (profile_name, budget.info.client, KIND_NAMES.get(budget.kind, ""), budget.info.dates) if x
    ]
    text(ws, 2, 1, " · ".join(info))
    r = 4
    heads(ws, r, ["Code", "Category", "Subtotal", "Fringes", "Total"])
    for section in SECTIONS:
        cats = [c for c in budget.categories if c.section == section]
        if not cats:
            continue
        r += 1
        text(ws, r, 2, SECTION_NAMES[section], bold)
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = band
        for c in cats:
            r += 1
            s = t.categories[c.id]
            text(ws, r, 1, c.code)
            text(ws, r, 2, c.name)
            num(ws, r, 3, s.subtotal)
            num(ws, r, 4, s.fringes)
            num(ws, r, 5, s.total)
        r += 1
        text(ws, r, 2, f"Total {SECTION_NAMES[section]}", bold)
        num(ws, r, 5, t.sections[section].total, bold)
    r += 2
    for label, value in (("Total Above the Line", t.atl.total), ("Total Below the Line", t.btl.total),
                         ("Above and Below the Line", t.direct.total)):  # fmt: skip
        text(ws, r, 2, label, bold)
        num(ws, r, 5, value, bold)
        r += 1
    for a, adj in zip(t.adjustments, budget.adjustments, strict=False):
        label = a.name or "Adjustment"
        if adj.kind == "percent":
            label += f" ({money.fmt_qty(money.parse(adj.rate))}% of {BASE_NAMES[adj.base]})"
        text(ws, r, 2, label)
        num(ws, r, 5, a.amount)
        r += 1
    text(ws, r, 2, "Grand Total", Font(bold=True, size=13))
    num(ws, r, 5, t.grand, Font(bold=True, size=13))
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 44
    for col in "CDE":
        ws.column_dimensions[col].width = 16

    # ---- detail
    ws = wb.create_sheet("Detail")
    heads(ws, 1, ["Code", "Description", "Amount", "Units", "Unit", "Rate", "Subtotal", "Fringes", "Total"])
    r = 1
    for c in budget.categories:
        r += 1
        text(ws, r, 1, c.code, bold)
        text(ws, r, 2, c.name, bold)
        num(ws, r, 9, t.categories[c.id].total, bold)
        for col in range(1, 10):
            ws.cell(row=r, column=col).fill = band
        for a in c.accounts:
            r += 1
            text(ws, r, 1, a.code, bold)
            text(ws, r, 2, a.name, bold)
            num(ws, r, 9, t.accounts[a.id].total, bold)
            for ln in a.lines:
                r += 1
                lt = t.lines[ln.id]
                text(ws, r, 2, "    " + ln.description)
                num(ws, r, 3, lt.amount, money_fmt=False)
                num(ws, r, 4, lt.units, money_fmt=False)
                text(ws, r, 5, UNIT_NAMES[ln.unit])
                num(ws, r, 6, lt.rate)
                num(ws, r, 7, lt.subtotal)
                num(ws, r, 8, lt.fringes)
                num(ws, r, 9, lt.total)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 44
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 13

    # ---- globals and fringes
    ws = wb.create_sheet("Globals")
    heads(ws, 1, ["Name", "What It Is", "Value", "Comes To"])
    for i, g in enumerate(budget.globals, 2):
        text(ws, i, 1, g.name)
        text(ws, i, 2, g.label)
        text(ws, i, 3, g.value)
        v = t.globals.get(g.name.casefold())
        if v is not None:
            num(ws, i, 4, v, money_fmt=False)
    ws = wb.create_sheet("Fringes")
    heads(ws, 1, ["Fringe", "Kind", "Rate", "Cap", "Lines", "Total"])
    for i, f in enumerate(budget.fringes, 2):
        text(ws, i, 1, f.name)
        text(ws, i, 2, "Percent" if f.kind == "percent" else "Flat per Unit")
        text(ws, i, 3, f.rate)
        text(ws, i, 4, f.cap)
        ws.cell(row=i, column=5, value=t.fringe_lines.get(f.id, 0))
        num(ws, i, 6, t.fringes.get(f.id, Decimal(0)))
    for ws in wb.worksheets[2:]:
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 30

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_name(f".{dest.name}.part")
    try:
        wb.save(tmp)
        os.replace(tmp, dest)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)
