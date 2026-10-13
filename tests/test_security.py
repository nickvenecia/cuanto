"""The protections: links, downloads, spreadsheets, files from other tools, text shown on screen."""

from __future__ import annotations

import zlib

import pytest

from cuanto.core import exchange, sheets, templates, web
from cuanto.core.model import Budget, Rate


def test_only_web_links_count():
    assert web.web_link("example.com/rates/1") == "https://example.com/rates/1"
    assert web.web_link("https://example.com:8443/x") == "https://example.com:8443/x"
    for bad in (
        "file:///etc/passwd",
        "file://attacker.com/share/x.exe",
        "smb://server/share",
        "javascript:alert(1)",
        "mailto:a@b.com",
        "ms-settings:display",
        "localhost:8080/x",
        "https://exa mple.com",
        "",
    ):
        assert web.web_link(bad) == "", bad


@pytest.mark.parametrize(
    "host_ip", ["127.0.0.1", "10.0.0.1", "192.168.1.1", "169.254.169.254", "::1", "fd00::1"]
)
def test_private_addresses_are_blocked(monkeypatch, host_ip):
    import socket

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", (host_ip, 443))])
    with pytest.raises(web.Blocked):
        web.check_public("router.example.com")


def test_public_addresses_pass(monkeypatch):
    import socket

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("93.184.215.14", 443))])
    web.check_public("example.com")


def test_a_redirect_to_a_private_address_is_refused(monkeypatch):
    import socket
    from email.message import Message

    def resolve(host, *a, **k):
        return [(0, 0, 0, "", ("93.184.215.14" if host == "shop.example.com" else "192.168.1.1", 443))]

    def get(url, ip, timeout, max_bytes):
        h = Message()
        h["Location"] = "https://router.local.example/admin"
        return web._Reply(302, h, b"")

    monkeypatch.setattr(socket, "getaddrinfo", resolve)
    monkeypatch.setattr(web, "_get", get)
    with pytest.raises(web.Blocked):
        web.fetch("https://shop.example.com/item")


def test_it_connects_to_the_address_it_checked(monkeypatch):
    import socket
    from email.message import Message

    seen = []
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("93.184.215.14", 443))])

    def get(url, ip, timeout, max_bytes):
        seen.append(ip)
        h = Message()
        h["Content-Type"] = "text/html; charset=idna"  # a charset Python can't decode with
        return web._Reply(200, h, b"<title>ok</title>")

    monkeypatch.setattr(web, "_get", get)
    page = web.fetch("https://shop.example.com/x")
    assert seen == ["93.184.215.14"] and "ok" in page.html


def test_broken_redirects_are_refused_politely(monkeypatch):
    import socket
    from email.message import Message

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("93.184.215.14", 443))])

    def get(url, ip, timeout, max_bytes):
        h = Message()
        h["Location"] = "https://[oops/"
        return web._Reply(302, h, b"")

    monkeypatch.setattr(web, "_get", get)
    with pytest.raises(OSError):
        web.fetch("https://shop.example.com/x")


@pytest.mark.parametrize("status", [401, 403, 429, 503])
def test_a_site_that_turns_cuanto_away_is_named(monkeypatch, status):
    import socket
    from email.message import Message

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("93.184.215.14", 443))])
    monkeypatch.setattr(web, "_get", lambda url, ip, timeout, max_bytes: web._Reply(status, Message(), b""))
    with pytest.raises(web.Refused, match=r"^rates\.example\.com "):
        web.fetch("https://www.rates.example.com/today.json")


def test_cuanto_says_what_it_is():
    assert web.HEADERS["User-Agent"].startswith("Cuanto/")


def test_tricky_links_are_not_web_links():
    for bad in (
        "https://[x",
        "https://[::1",
        "http://[evil.com]/",
        "https://example.com@evil.com/x",
        "https://a:b@evil.com",
        "https://evil.com:99999/",
        "https://evil.com:abc/",
        "https://ev%69l.com",
    ):
        assert web.web_link(bad) == "", bad


@pytest.mark.parametrize("addr", ["::ffff:127.0.0.1", "64:ff9b::7f00:1", "::127.0.0.1", "2002:c0a8:0101::1"])
def test_ipv6_wrappers_of_private_addresses_are_blocked(monkeypatch, addr):
    import socket

    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", (addr, 443))])
    with pytest.raises(web.Blocked):
        web.check_public("sneaky.example.com")


def test_unpacking_is_capped():
    bomb = zlib.compress(b"0" * 50_000_000)  # 50 MB of zeros packs to ~50 KB
    out = web._unpack(bomb, "deflate", 1_000_000)
    assert len(out) <= 1_000_000


def test_an_excel_export_cant_carry_formulas(tmp_path):
    from openpyxl import load_workbook

    b = templates.build("short", '=HYPERLINK("http://evil.example","x")')
    b.categories[0].accounts[0].lines[0].description = "=1+2"
    out = tmp_path / "evil.xlsx"
    sheets.write_xlsx(b, out)
    wb = load_workbook(out)
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any(c.data_type == "f" for c in cells)
    assert any(c.value == "=1+2" or c.value == "    =1+2" for c in cells)


def test_a_csv_too_big_for_a_budget_is_refused(tmp_path, monkeypatch):
    path = tmp_path / "huge.csv"
    path.write_text("Description,Rate\n" + "x,1\n" * 100)
    monkeypatch.setattr(sheets, "MAX_BYTES", 100)
    with pytest.raises(sheets.SheetError):
        sheets.read_csv(path)


def test_json_from_another_tool_is_only_data(tmp_path):
    evil = {
        "format": "cuanto.exchange",
        "version": 1,
        "rates": [
            {"name": "x" * 50_000, "rate": "__import__('os')", "union": "yes", "fringes": [{"a": 1}, "FICA"]}
        ],
        "schedule": {"globals": {"shoot days": "1e400", "": "3"}},
        "crew": [{"role": {"nested": True}, "name": "", "rate": "5"}, "not a record"],
    }
    lib: dict[str, Rate] = {}
    b = Budget()
    rep = exchange.apply(exchange.check(evil), lib, b)
    r = next(iter(lib.values()))
    assert len(r.name) <= 2000 and r.union is True and r.fringes == ["FICA"]
    assert rep.rates_added == 1 and rep.globals_set == 1 and rep.lines_added == 0
    assert b.globals[0].name == "shoot_days"  # a name lines can use, whatever came in


def test_a_file_from_a_newer_cuanto_is_refused_with_words(tmp_path):
    with pytest.raises(exchange.ExchangeError, match="newer"):
        exchange.check({"format": "cuanto.exchange", "version": 99})
    with pytest.raises(exchange.ExchangeError):
        exchange.check({"format": "something.else"})
    deep = tmp_path / "deep.json"
    deep.write_text("[" * 100_000)
    with pytest.raises(exchange.ExchangeError):
        exchange.read(deep)


def test_names_show_as_typed(qapp):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QLabel

    from cuanto.app import setup_app

    setup_app(qapp)
    lbl = QLabel('<img src="x.png"> Evil')
    lbl.ensurePolished()
    assert lbl.textFormat() == Qt.TextFormat.PlainText
