"""The window, offscreen: making a budget, editing lines in the panel and in the table, fringes, globals,
deleting, undo, importing and exporting, the page window, and what's dropped when left empty."""

from __future__ import annotations

import json
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path

import pytest
from PySide6.QtCore import Qt, QThread, QTimer
from PySide6.QtGui import QContextMenuEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QWidget

from cuanto.core import calc, sample
from cuanto.core.store import Library


@pytest.fixture
def win(qapp, data_folder):
    from cuanto.ui.window.main_window import MainWindow

    w = MainWindow(sample.fill(Library(data_folder).load()))
    w._restored = True
    w.resize(1600, 1000)
    w.show()
    yield w
    w.close()


def feature(w):
    return next(b for b in w.lib.budgets.values() if b.kind == "feature")


def say_yes(monkeypatch) -> None:
    monkeypatch.setattr("cuanto.ui.window.editing.confirm", lambda *a, **k: True)


def when_open(kind, do) -> None:
    """`do(window)` once a window of this kind shows (it may open from inside a click and wait there)."""
    tries = [500]

    def look() -> None:
        shown = [w for w in QApplication.topLevelWidgets() if isinstance(w, kind) and w.isVisible()]
        if shown:
            do(shown[0])
        elif tries[0]:
            tries[0] -= 1
            QTimer.singleShot(10, look)

    QTimer.singleShot(0, look)


def wait_for(done, ms: int = 5000) -> None:
    for _ in range(ms // 10):
        if done():
            return
        QTest.qWait(10)
    raise AssertionError("never happened")


@contextmanager
def still_there_after(widget: QWidget):
    """A click on `widget` may open a window or a menu; anything that window throws away while it's open is
    gone before the click is over. If that's `widget` itself, Cuánto crashes (on a Mac, at once). Fails if it
    happens during the `with`."""
    start = QThread.currentThread().loopLevel()
    thrown = []
    original = QWidget.deleteLater

    def delete_later(w: QWidget) -> None:
        if QThread.currentThread().loopLevel() > start and (w is widget or w.isAncestorOf(widget)):
            thrown.append(type(widget).__name__)
        original(w)

    QWidget.deleteLater = delete_later
    try:
        yield
    finally:
        QWidget.deleteLater = original
    assert not thrown, f"the {thrown[0]} was thrown away in the middle of its own click"


def test_first_launch_shows_the_ways_to_start(qapp, data_folder):
    from cuanto.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    assert w.pages.currentWidget() is w.empty and len(w.empty.tiles) == 5
    assert not w.search.isVisibleTo(w) and not w.sidebar.heads["open"].isVisibleTo(w.sidebar)
    w.close()


def test_saving_your_profile_while_typing_in_it(qapp, data_folder, monkeypatch):
    """Profile card › Edit Profile…, type a name, Save. Saving puts a fresh form in place of the one you typed
    in, and throwing that one away with the typing still in it crashed Cuánto, every time."""
    from PySide6.QtWidgets import QMenu, QPushButton

    from cuanto.ui.dialogs.settings import SettingsDialog
    from cuanto.ui.window.main_window import MainWindow

    monkeypatch.setattr("cuanto.ui.dialogs.common.confirm", lambda *a, **k: True)
    w = MainWindow(Library(data_folder).load())
    w._restored = True
    w.show()
    card = w.sidebar.profile_card

    def one_by_one(*steps) -> None:
        """Each step once Settings has had a moment to itself, as between a person's clicks."""
        if steps:
            QTimer.singleShot(50, lambda: (steps[0](), one_by_one(*steps[1:])))

    def type_and_click(settings, field: str, text: str, button: str = "") -> None:
        form = settings.profile_editor
        getattr(form, field).setFocus()
        getattr(form, field).selectAll()
        QTest.keyClicks(getattr(form, field), text)
        click = (
            form.save if not button else next(b for b in form.findChildren(QPushButton) if b.text() == button)
        )
        QTest.mouseClick(click, Qt.MouseButton.LeftButton)

    def edit(settings) -> None:
        settings.activateWindow()
        assert QTest.qWaitForWindowActive(settings)  # (the typing only goes to the window in front)
        one_by_one(
            lambda: type_and_click(settings, "name", "Juniper Lane Films"),  # Add Profile
            lambda: type_and_click(settings, "email", "hello@example.com"),  # Save Changes
            settings.close,
        )

    def pick(menu) -> None:
        when_open(SettingsDialog, edit)
        next(a for a in menu.actions() if a.text() == "Edit Profile…").trigger()
        menu.close()

    when_open(QMenu, pick)
    with still_there_after(card):
        QTest.mouseClick(card, Qt.MouseButton.LeftButton)
    saved = list(Library(data_folder).load().profiles.values())
    assert [(p.name, p.email) for p in saved] == [("Juniper Lane Films", "hello@example.com")]
    assert card.name.full_text() == "Juniper Lane Films"

    def delete(settings) -> None:
        settings.activateWindow()
        assert QTest.qWaitForWindowActive(settings)
        one_by_one(lambda: type_and_click(settings, "name", "Juniper", "Delete Profile…"), settings.close)

    when_open(SettingsDialog, delete)
    w.edit_profile()  # Settings › Profiles straight away
    assert Library(data_folder).load().profiles == {}
    w.close()


def test_deleting_a_budget_from_its_row_in_the_sidebar(win, monkeypatch):
    """Right-click › Delete… takes away the very row that was right-clicked, so its menu waits until the
    right-click is over."""
    from PySide6.QtWidgets import QMenu

    say_yes(monkeypatch)
    b = feature(win)
    row = win.sidebar.rows[f"budget:{b.id}"]

    def delete(menu) -> None:
        next(a for a in menu.actions() if a.text() == "Delete…").trigger()
        menu.close()

    when_open(QMenu, delete)
    at = row.rect().center()
    with still_there_after(row):
        QApplication.sendEvent(
            row, QContextMenuEvent(QContextMenuEvent.Reason.Mouse, at, row.mapToGlobal(at))
        )
    wait_for(lambda: b.id not in win.lib.budgets)
    assert b.id not in Library(win.lib.root).load().budgets
    assert f"budget:{b.id}" not in win.sidebar.rows


def test_the_mac_title_bar_never_fits_a_window_that_is_gone(qapp, data_folder, monkeypatch):
    """The Mac's title bar fits itself just after the window shows. A window thrown away before then is left
    alone: reaching into it would break every test after it on the Mac."""
    from PySide6.QtCore import QCoreApplication, QEvent

    from cuanto.ui.components import titlebar
    from cuanto.ui.window.main_window import MainWindow

    monkeypatch.setattr(titlebar, "supported", lambda: True)  # the Mac's bar, on any computer
    fits = []
    monkeypatch.setattr(titlebar.TitleBar, "fit", lambda bar: fits.append(bar))
    w = MainWindow(Library(data_folder).load())
    w._restored = True
    w.show()
    wait_for(lambda: fits)  # a window that stays fits its bar
    gone = MainWindow(Library(data_folder).load())
    gone._restored = True
    gone.show()
    gone.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    QTest.qWait(20)
    assert len(fits) == 1
    w.close()


def test_a_new_budget_from_a_template(qapp, data_folder, monkeypatch):
    from cuanto.ui.dialogs import pickers
    from cuanto.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    monkeypatch.setattr(pickers.NewBudgetDialog, "exec", lambda self: (self.field.setText("Spot"), 1)[1])
    monkeypatch.setattr(pickers.NewBudgetDialog, "kind", property(lambda self: "commercial"))
    w.new_budget()
    assert w.budget is not None and w.budget.name == "Spot" and w.view_key == "top"
    assert w.sidebar.heads["open"].label.full_text() == "Spot"
    assert Library(data_folder).load().budgets[w.budget.id].kind == "commercial"
    w.close()


def test_the_wheel_scrolls_the_panel_never_a_title_or_the_typing(win):
    """A long name wraps onto more lines and shows whole: turning the wheel over it scrolls the panel, never the
    name's own words out of sight. And scrolling past a dropdown or a date never takes the typing."""
    from PySide6.QtCore import QPoint, QPointF
    from PySide6.QtGui import QWheelEvent
    from PySide6.QtWidgets import QAbstractSpinBox, QComboBox

    b = feature(win)
    win.go(f"budget:{b.id}")
    win.open("line", b.categories[0].accounts[0].lines[0].id)
    title = win.line_panel.name
    for words in (1, 4, 12, 30):
        title.setText(" ".join(["Steadicam"] * words))
        QTest.qWait(20)
        assert title.isVisible() and title.verticalScrollBar().maximum() == 0, f"{words} words scroll"
        at = QPointF(10, 10)
        down = QWheelEvent(at, QPointF(title.mapToGlobal(QPoint(10, 10))), QPoint(), QPoint(0, -120),
                           Qt.MouseButton.NoButton, Qt.KeyboardModifier.NoModifier, Qt.ScrollPhase.NoScrollPhase,
                           False)  # fmt: skip
        QApplication.sendEvent(title.viewport(), down)
        assert title.verticalScrollBar().value() == 0 and not down.isAccepted()  # it goes on to the panel

    controls = [
        c for p in win.panels.values() for t in (QComboBox, QAbstractSpinBox) for c in p.findChildren(t)
    ]
    assert any(isinstance(c, QComboBox) for c in controls)
    for c in controls:
        c.ensurePolished()
        assert c.focusPolicy() == Qt.FocusPolicy.StrongFocus  # a click or Tab, never the wheel going past


def test_typing_a_line_in_its_panel_saves_and_adds_up(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    camera = next(c for c in b.categories if c.name == "Camera")
    acc = camera.accounts[0]
    win.new_line(account_id=acc.id)
    panel = win.line_panel
    ln = panel.ln
    assert win.view_key == "sec:production" and ln in acc.lines
    panel.name.setPlainText("Steadicam Operator")  # typing
    panel.units.setText("shoot_days")
    panel.units.textEdited.emit("shoot_days")
    panel.rate.setText("1200")
    panel.rate.textEdited.emit("1200")
    fica = b.fringes[0]
    panel.fringe_rows[fica.id].setChecked(True)
    win.flush()
    again = Library(win.lib.root).load().budgets[b.id]
    got = again.find_line(ln.id)[2]
    assert got.description == "Steadicam Operator" and got.units == "shoot_days" and got.rate == "1200"
    assert got.fringes == [fica.id]
    t = calc.totals(again)
    assert t.lines[ln.id].subtotal == Decimal("24000.00") and t.lines[ln.id].fringes == Decimal("1836.00")
    assert panel.total.text() == "$25,836.00"


def test_typing_over_a_line_in_the_table(win):
    b = feature(win)
    win.go("sec:production")
    win.go(f"budget:{b.id}")
    win.go("sec:production")
    ln = next(ln for _c, _a, ln in b.lines() if ln.description == "Grips")
    win.cell_edited(ln.id, "units", "shoot_days + 5")
    win.cell_edited(ln.id, "unit", "weeks")
    got = Library(win.lib.root).load().budgets[b.id].find_line(ln.id)[2]
    assert got.units == "shoot_days + 5" and got.unit == "weeks"
    row = win.tables["section"].find("line", ln.id)
    assert row is not None and row.text(win.tables["section"].col("units")) == "25"


def test_changing_a_global_changes_every_line(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    before = calc.totals(b).grand
    g = b.global_named("shoot_days")
    win.open("global", g.id)
    win.global_panel.value.setText("25")
    win.global_panel.value.textEdited.emit("25")
    win.flush()
    win.refresh()
    assert calc.totals(Library(win.lib.root).load().budgets[b.id]).grand > before
    win.global_panel.name.setText("days_shooting")
    win.global_panel.name.editingFinished.emit()
    assert g.name == "days_shooting" and calc.uses_of(b, "days_shooting") > 10
    assert not calc.totals(b).problems()


def test_an_account_switch_puts_a_fringe_on_every_line(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    acc = next(a for _c, a in b.accounts() if a.name == "Camera Equipment")
    union = next(f for f in b.fringes if f.name == "Union P&H")
    win.open("account", acc.id)
    win.account_panel.fringe_rows[union.id].setChecked(True)
    assert all(union.id in ln.fringes for ln in acc.lines)
    win.account_panel.fringe_rows[union.id].setChecked(False)
    assert not any(union.id in ln.fringes for ln in acc.lines)


def test_a_new_line_left_empty_is_not_kept(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    before = sum(1 for _ in b.lines())
    win.new_line()
    assert sum(1 for _ in b.lines()) == before + 1
    win.deselect()  # Esc
    assert sum(1 for _ in b.lines()) == before
    assert sum(1 for _ in Library(win.lib.root).load().budgets[b.id].lines()) == before


def test_deleting_asks_first_and_undo_brings_it_back(win, monkeypatch):
    b = feature(win)
    win.go(f"budget:{b.id}")
    cat = next(c for c in b.categories if c.name == "Camera")
    asked = []
    monkeypatch.setattr(
        "cuanto.ui.window.editing.confirm",
        lambda _w, title, text, *a, **k: asked.append((title, text)) or True,
    )
    win.delete_many([("category", cat.id)])
    assert asked == [("Delete 2700 Camera?", "Its 2 accounts and 6 lines go too.")]
    assert win.budget.find_category(cat.id) is None
    win.undo()
    assert win.budget.find_category(cat.id) is not None
    assert Library(win.lib.root).load().budgets[b.id].find_category(cat.id) is not None


def test_deleting_a_fringe_takes_it_off_every_line(win, monkeypatch):
    say_yes(monkeypatch)
    b = feature(win)
    win.go(f"budget:{b.id}")
    fica = b.fringes[0]
    win.delete_many([("fringe", fica.id)])
    assert not any(fica.id in ln.fringes for _c, _a, ln in b.lines())


def test_moving_lines_and_categories(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    acc = next(a for _c, a in b.accounts() if a.name == "Camera Crew")
    first, second = acc.lines[0], acc.lines[1]
    win.step("line", first.id, 1)
    assert acc.lines[0] is second and acc.lines[1] is first
    atl = [c for c in b.categories if c.section == "atl"]
    win.step("category", atl[0].id, 1)
    assert [c for c in b.categories if c.section == "atl"][1] is atl[0]


def test_rates_go_onto_lines_and_lines_into_the_library(win, monkeypatch):
    from cuanto.ui.dialogs import pickers

    b = feature(win)
    win.go(f"budget:{b.id}")
    ln = next(x for _c, _a, x in b.lines() if x.description == "Gaffer")
    gaffer = next(r for r in win.lib.rates.values() if r.name == "Gaffer" and r.union)
    monkeypatch.setattr(pickers.RatePicker, "exec", lambda self: 1)
    monkeypatch.setattr(pickers.RatePicker, "rate_id", property(lambda self: gaffer.id))
    win.use_rate(ln.id)
    assert ln.rate == "725" and ln.rate_id == gaffer.id
    assert {b.fringe(f).name for f in ln.fringes} == {"FICA", "Union P&H"}
    n = len(win.lib.rates)
    dit = next(x for _c, _a, x in b.lines() if x.description == "DIT")
    win.save_rate(dit.id)
    assert len(win.lib.rates) == n + 1 and dit.rate_id in win.lib.rates


def test_import_csv_into_the_open_budget_and_json_from_another_tool(win, tmp_path):
    b = feature(win)
    win.go(f"budget:{b.id}")
    win.go("sec:production")
    sheet = tmp_path / "extra.csv"
    sheet.write_text(
        "Category,Category Code,Account,Description,Units,Unit,Rate\nCamera,2700,Camera Crew,Steadicam,5,days,1100\n"
    )
    win.import_file(sheet)
    assert any(ln.description == "Steadicam" for _c, _a, ln in b.lines())
    assert list((win.lib.root / "backups").glob(f"{b.id}-*.json"))  # a copy was kept first
    crew = tmp_path / "crew.json"
    crew.write_text(json.dumps({"format": "cuanto.exchange", "version": 1,
                                "schedule": {"globals": {"shoot_days": 30}},
                                "crew": [{"role": "Gaffer", "name": "A. Rivera", "account": "2901", "rate": "700"}]}))  # fmt: skip
    win.import_file(crew)
    assert b.global_named("shoot_days").value == "30"
    assert any(ln.description == "Gaffer — A. Rivera" for _c, _a, ln in b.lines())


def test_exports_and_the_page_window(win, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from cuanto.ui.dialogs.preview import PageDialog

    b = feature(win)
    win.go(f"budget:{b.id}")
    out = iter([tmp_path / "b.xlsx", tmp_path / "b.json", tmp_path / "b.pdf"])
    monkeypatch.setattr(
        "cuanto.ui.window.transfer.QFileDialog.getSaveFileName", lambda *a, **k: (str(next(out)), "")
    )
    win.export_xlsx()
    assert load_workbook(tmp_path / "b.xlsx").sheetnames[0] == "Top Sheet"
    win.export_json()
    assert json.loads((tmp_path / "b.json").read_text())["format"] == "cuanto.budget"
    shown = []

    def page(d):
        shown.append(d.panel.view.pages)
        d.panel.save_pdf.emit()
        return 0

    monkeypatch.setattr(PageDialog, "exec", page)
    win.show_pages("both")
    assert shown and len(shown[0]) >= 3
    assert (tmp_path / "b.pdf").read_bytes().startswith(b"%PDF")


def test_search_and_the_lines_filter(win):
    from cuanto.ui.window.tables import ID, KIND

    b = feature(win)
    win.go(f"budget:{b.id}")
    win.go("sec:production")
    t = win.tables["section"]
    win.search.setText("gaffer")
    shown = [r.data(0, ID) for r in t.showing_rows() if r.data(0, KIND) == "line"]
    assert len(shown) == 1
    win.search.clear()
    ln = next(x for _c, _a, x in b.lines() if x.description == "Grips")
    ln.rate = ""
    win.refresh()
    win.filter_btn._pick("show", "no_rate")
    lines = [r.data(0, ID) for r in t.showing_rows() if r.data(0, KIND) == "line"]
    assert ln.id in lines and all(not b.find_line(i)[2].rate for i in lines)
    win.filter_btn._clear()


def test_a_problem_shows_in_the_table_and_the_panel(win):
    b = feature(win)
    win.go(f"budget:{b.id}")
    win.go("sec:production")
    ln = next(x for _c, _a, x in b.lines() if x.description == "Grips")
    win.cell_edited(ln.id, "units", "shoot_dayz")
    row = win.tables["section"].find("line", ln.id)
    assert row.text(win.tables["section"].col("total")) == "—"
    win.open("line", ln.id)
    assert (
        win.line_panel.problem.isVisibleTo(win.line_panel) and "shoot_dayz" in win.line_panel.problem.text()
    )


def test_deleting_the_open_budget_goes_back_to_all_budgets(win, monkeypatch):
    say_yes(monkeypatch)
    b = feature(win)
    win.go(f"budget:{b.id}")
    win.delete_many([("budget", b.id)])
    assert win.budget is None and win.view_key == "budgets" and b.id not in win.lib.budgets
    assert not (Path(win.lib.root) / "budgets" / f"{b.id}.json").exists()


def test_a_budget_saved_as_a_template_starts_the_next_one(win, monkeypatch):
    from cuanto.ui.dialogs import pickers

    b = feature(win)
    win.go(f"budget:{b.id}")
    monkeypatch.setattr(pickers.SaveTemplateDialog, "exec", lambda self: (self.field.setText("Ours"), 1)[1])
    win.save_as_template()
    t = win.lib.template_named("Ours")
    assert t is not None and Library(win.lib.root).load().templates[t.id].budget.kind == "feature"

    seen = {}

    def pick_ours(self):
        seen["choices"] = [self.template.itemData(i) for i in range(self.template.count())]
        self.field.setText("The Next Night")
        self.template.setCurrentIndex(self.template.findData("mine:" + t.id))
        return 1

    monkeypatch.setattr(pickers.NewBudgetDialog, "exec", pick_ours)
    win.new_budget()
    assert "feature" in seen["choices"] and "mine:" + t.id in seen["choices"]
    new = win.budget
    assert new.name == "The Next Night" and new.id != b.id
    assert calc.totals(new).grand == calc.totals(b).grand
    assert win.lib.settings["new_from"] == "mine:" + t.id  # offered first next time

    monkeypatch.setattr("cuanto.ui.dialogs.settings.confirm", lambda *a, **k: True)
    from cuanto.ui.dialogs.settings import SettingsDialog

    d = SettingsDialog(win.lib, win, "templates")
    d.panes["templates"].findChild(pickers.QPushButton, "Danger").click()
    assert win.lib.templates == {} and Library(win.lib.root).load().templates == {}
    win.undo()
    assert t.id in win.lib.templates
