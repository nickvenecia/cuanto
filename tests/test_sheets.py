"""Spreadsheets (core/sheets.py): budget lines in from CSV, a budget out to Excel."""

from __future__ import annotations

from decimal import Decimal

import pytest

from cuanto.core import calc, sample, sheets, templates
from cuanto.core.store import Library


def _csv(tmp_path, text: str):
    path = tmp_path / "lines.csv"
    path.write_text(text, encoding="utf-8")
    return path


def test_a_sheet_with_categories_comes_in(tmp_path):
    path = _csv(tmp_path, (
        "My Budget,,,,,,\n"
        ",,,,,,\n"
        "Cat #,Category,Account,Description,Qty,Units,Unit,Rate,Fringes\n"
        "1300,Director,Director,Director Fee,1,1,Flat,\"$50,000\",FICA\n"
        "2700,Camera,Crew,DP,1,20,days,900,FICA; Union P&H\n"
        "2700,Camera,Equipment,Camera Package,1,4,wks,7500,\n"
        "4400,Color,Color,Grade,,,,,\n"
    ))  # fmt: skip
    got = sheets.read_csv(path)
    assert got.lines == 4 and got.skipped == 0  # a line with no rate yet is still a line
    assert [c.section for c in got.categories] == ["atl", "production", "post"]
    camera = got.categories[1]
    assert [a.name for a in camera.accounts] == ["Crew", "Equipment"]
    pkg = camera.accounts[1].lines[0]
    assert pkg.unit == "weeks" and pkg.units == "4" and pkg.rate == "7500"
    b = templates.build("blank", "Imported")
    b.fringes = templates.build("feature").fringes
    assert sheets.merge_into(b, got) == 4
    assert calc.totals(b).direct.subtotal == Decimal(50000 + 18000 + 30000)
    dp = b.categories[1].accounts[0].lines[0]
    assert [b.fringe(f).name for f in dp.fringes] == ["FICA", "Union P&H"]


def test_headings_make_accounts_and_totals_become_allowances(tmp_path):
    path = _csv(tmp_path, (
        "Description;Total\n"
        "Locations;\n"
        "Site fees;1500\n"
        "Permits;(200)\n"
        "Catering;\n"
        "Lunch;2,400.50\n"
    ))  # fmt: skip
    got = sheets.read_csv(path)
    accounts = got.categories[0].accounts
    assert [a.name for a in accounts] == ["Locations", "Catering"]
    assert accounts[0].lines[1].rate == "-200" and accounts[0].lines[1].unit == "allow"
    assert accounts[1].lines[0].rate == "2400.50"


def test_a_sheet_without_column_titles_says_so(tmp_path):
    with pytest.raises(sheets.SheetError, match="column titles"):
        sheets.read_csv(_csv(tmp_path, "a,b,c\n1,2,3\n"))


def test_merging_joins_the_same_category(data_folder, tmp_path):
    lib = sample.fill(Library(data_folder).load())
    b = next(x for x in lib.budgets.values() if x.kind == "feature")
    before = len(b.categories)
    path = _csv(
        tmp_path,
        "Category,Category Code,Account,Description,Rate\nCamera,2700,Camera Crew,Steadicam Op,1100\n",
    )
    sheets.merge_into(b, sheets.read_csv(path))
    assert len(b.categories) == before
    crew = next(a for _c, a in b.accounts() if a.code == "2701")
    assert crew.lines[-1].description == "Steadicam Op"


def test_excel_has_every_sheet_and_the_grand_total(data_folder, tmp_path):
    from openpyxl import load_workbook

    lib = sample.fill(Library(data_folder).load())
    b = lib.ordered_budgets()[1]
    out = tmp_path / "budget.xlsx"
    sheets.write_xlsx(b, out, "Mesa Alta Pictures")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Top Sheet", "Detail", "Globals", "Fringes"]
    values = [c.value for row in wb["Top Sheet"].iter_rows() for c in row]
    assert round(Decimal(str(values[values.index("Grand Total") + 3])), 2) == calc.totals(b).grand
    assert not list(tmp_path.glob(".*.part"))
